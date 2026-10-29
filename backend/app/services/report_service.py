import os
import csv
import json
from datetime import datetime
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPORTS_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "data", "reports")

def ensure_reports_dir():
    os.makedirs(REPORTS_DIR, exist_ok=True)

def generate_pdf_report(incident: dict, detection: dict, validation: dict, drift: dict, vessels: list, evidence: list) -> str:
    ensure_reports_dir()
    inc_id = incident["id"]
    filename = f"SeaTrace_Investigation_Dossier_{inc_id.replace('-', '_')}.pdf"
    file_path = os.path.join(REPORTS_DIR, filename)

    doc = SimpleDocTemplate(
        file_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        textColor=colors.HexColor('#0A1020'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        textColor=colors.HexColor('#FF7A2F'),
        spaceAfter=12
    )
    heading2_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        textColor=colors.HexColor('#334155'),
        leading=13
    )

    story = []

    # Header
    story.append(Paragraph("SEATRACE AI — OFFICIAL MARITIME INCIDENT DOSSIER", title_style))
    story.append(Paragraph(f"CONFIDENTIAL INVESTIGATION DOSSIER • INCIDENT {inc_id} • GENERATED AT {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#FF7A2F'), spaceAfter=10))

    # Incident Overview Table
    story.append(Paragraph("1. Incident Summary & Geographic Coordinates", heading2_style))
    inc_table_data = [
        ["Incident ID", inc_id, "Classification", incident.get("severity", "N/A")],
        ["Title", incident.get("title", "N/A"), "Status", incident.get("status", "N/A")],
        ["Maritime Region", incident.get("region", "N/A"), "Coordinates", f"{incident.get('lat', 0):.4f}°N, {incident.get('lng', 0):.4f}°E"],
        ["Detection Time", incident.get("detection_time", "N/A"), "Estimated Volume", f"{incident.get('estimated_volume_m3', 0)} m³"],
        ["Sensor / Mission", incident.get("source", "N/A"), "Spill Type", incident.get("spill_type", "N/A")]
    ]
    t1 = Table(inc_table_data, colWidths=[110, 160, 110, 160])
    t1.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('PADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t1)
    story.append(Spacer(1, 10))

    # Satellite Detection & Screening
    story.append(Paragraph("2. Satellite Detection & Look-Alike Validation", heading2_style))
    if detection:
        det_text = f"Analyzed scene <b>{detection.get('scene_id', 'N/A')}</b> captured by {detection.get('satellite_mission', 'N/A')} {detection.get('sensor', 'N/A')}. Estimated slick footprint area is <b>{detection.get('area_km2', 0):.2f} km²</b> with perimeter {detection.get('perimeter_km', 0):.2f} km. Algorithmic confidence score is <b>{detection.get('confidence_score', 0)*100:.1f}%</b>."
    else:
        det_text = "No satellite detection record linked."
    story.append(Paragraph(det_text, body_style))
    story.append(Spacer(1, 4))

    if validation:
        val_text = f"<b>Validation Decision:</b> <font color='#16A34A'><b>{validation.get('decision', 'N/A')}</b></font> by reviewer {validation.get('reviewer_name', 'N/A')} ({validation.get('reviewer_role', 'N/A')}). Surface wind speed: {validation.get('wind_speed_knots', 0)} kts. Biogenic surfactant risk: {validation.get('biogenic_risk', 'N/A')}. Proximity to charted natural seep: {validation.get('natural_seep_proximity_km', 0)} km. Notes: <i>{validation.get('notes', 'None')}</i>"
        story.append(Paragraph(val_text, body_style))
    story.append(Spacer(1, 10))

    # Reverse Drift Reconstruction
    story.append(Paragraph("3. Environmental Origin Reconstruction", heading2_style))
    if drift:
        drift_text = f"Lagrangian backward trajectory simulation executed using <b>{drift.get('particle_count', 0)} particles</b> over a <b>{drift.get('time_window_hours', 0)}-hour window</b> with windage leeway {drift.get('windage_leeway', 0)*100:.1f}%. Estimated release origin centroid is located at <b>{drift.get('origin_lat', 0):.4f}°N, {drift.get('origin_lng', 0):.4f}°E</b> (spatial uncertainty ±{drift.get('origin_uncertainty_km', 0)} km). Forcing models: {drift.get('environmental_sources', 'N/A')}."
        story.append(Paragraph(drift_text, body_style))
    else:
        story.append(Paragraph("Drift reconstruction not yet executed.", body_style))
    story.append(Spacer(1, 10))

    # Candidate Vessel Ranking
    story.append(Paragraph("4. AIS Vessel Trajectory Correlation & Attribution", heading2_style))
    vessel_table_data = [["Rank", "MMSI", "Vessel Name", "Flag", "Type", "Dist (km)", "Time Δ", "Score", "Lead Status"]]
    for idx, v in enumerate(vessels[:5]):
        vessel_table_data.append([
            f"#{idx+1}",
            str(v.get("mmsi")),
            v.get("vessel_name", "Unknown"),
            v.get("flag", "Unknown"),
            v.get("vessel_type", "Unknown"),
            f"{v.get('distance_to_origin_km', 0):.1f}",
            f"{v.get('time_offset_hours', 0):+.1f}h",
            f"{v.get('compatibility_score', 0):.1f}%",
            v.get("risk_level", "N/A")
        ])
    if len(vessel_table_data) > 1:
        t2 = Table(vessel_table_data, colWidths=[35, 65, 120, 60, 85, 50, 40, 45, 60])
        t2.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7.5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('PADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t2)
    else:
        story.append(Paragraph("No candidate vessels correlated for this time window.", body_style))
    story.append(Spacer(1, 10))

    # Evidence Manifest
    story.append(Paragraph("5. Cryptographic Evidence Chain of Custody", heading2_style))
    ev_table_data = [["Artifact Type", "Title", "SHA-256 Checksum Hash", "Verified"]]
    for ev in evidence[:6]:
        ev_table_data.append([
            ev.get("artifact_type", ""),
            ev.get("title", ""),
            f"{ev.get('sha256_hash', '')[:16]}...{ev.get('sha256_hash', '')[-8:]}",
            "VERIFIED" if ev.get("verified") else "PENDING"
        ])
    if len(ev_table_data) > 1:
        t3 = Table(ev_table_data, colWidths=[120, 180, 180, 60])
        t3.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7.5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('PADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t3)
    story.append(Spacer(1, 15))

    # Legal Disclaimer / Signature
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#94A3B8'), spaceAfter=8))
    story.append(Paragraph("<b>LEGAL NOTICE:</b> This investigative dossier contains probabilistic modeling, satellite remote sensing interpretations, and AIS track correlation. Vessel rankings constitute candidate investigative leads for competent maritime enforcement authorities and do not constitute self-executing legal guilt without field corroboration.", ParagraphStyle('Notice', parent=body_style, fontSize=7.5, textColor=colors.HexColor('#64748B'))))

    doc.build(story)
    return file_path

def generate_csv_report(incident: dict, vessels: list) -> str:
    ensure_reports_dir()
    inc_id = incident["id"]
    filename = f"SeaTrace_Vessel_Correlation_{inc_id.replace('-', '_')}.csv"
    file_path = os.path.join(REPORTS_DIR, filename)

    with open(file_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Incident ID", inc_id])
        writer.writerow(["Title", incident.get("title", "")])
        writer.writerow(["Region", incident.get("region", "")])
        writer.writerow(["Coordinates", f"{incident.get('lat', '')}, {incident.get('lng', '')}"])
        writer.writerow([])
        writer.writerow([
            "MMSI", "Vessel Name", "Flag", "IMO", "Vessel Type",
            "Speed (knots)", "Course (deg)", "Dist to Origin (km)",
            "Time Offset (hours)", "Draught Delta (m)", "Compatibility Score (%)", "Risk Level"
        ])
        for v in vessels:
            writer.writerow([
                v.get("mmsi"),
                v.get("vessel_name"),
                v.get("flag"),
                v.get("imo", ""),
                v.get("vessel_type"),
                v.get("speed_knots"),
                v.get("course_deg"),
                v.get("distance_to_origin_km"),
                v.get("time_offset_hours"),
                v.get("draught_change_m"),
                v.get("compatibility_score"),
                v.get("risk_level")
            ])
    return file_path

def generate_xlsx_report(incident: dict, detection: dict, validation: dict, drift: dict, vessels: list, evidence: list) -> str:
    ensure_reports_dir()
    inc_id = incident["id"]
    filename = f"SeaTrace_Investigation_Workbook_{inc_id.replace('-', '_')}.xlsx"
    file_path = os.path.join(REPORTS_DIR, filename)

    wb = openpyxl.Workbook()
    # Sheet 1: Incident Summary
    ws1 = wb.active
    ws1.title = "Incident Summary"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0A1020", end_color="0A1020", fill_type="solid")
    title_font = Font(name="Segoe UI", size=14, bold=True, color="0A1020")

    ws1.append(["SEATRACE AI — MARITIME INVESTIGATION DOSSIER"])
    ws1["A1"].font = title_font
    ws1.append(["Incident Reference", inc_id])
    ws1.append(["Title", incident.get("title")])
    ws1.append(["Region", incident.get("region")])
    ws1.append(["Status", incident.get("status")])
    ws1.append(["Severity", incident.get("severity")])
    ws1.append(["Latitude", incident.get("lat")])
    ws1.append(["Longitude", incident.get("lng")])
    ws1.append(["Detection Timestamp", incident.get("detection_time")])
    ws1.append(["Source Sensor", incident.get("source")])
    ws1.append(["Estimated Volume (m³)", incident.get("estimated_volume_m3")])
    ws1.append(["Slick Area (km²)", detection.get("area_km2") if detection else "N/A"])
    ws1.append(["Validation Review", validation.get("decision") if validation else "UNREVIEWED"])
    ws1.append(["Reviewer", validation.get("reviewer_name") if validation else "N/A"])

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Sheet 2: Vessel Candidates
    ws2 = wb.create_sheet(title="Vessel Candidates")
    headers2 = ["MMSI", "Vessel Name", "Flag", "IMO", "Vessel Type", "Speed (kts)", "Course (°)", "Dist to Origin (km)", "Time Δ (hrs)", "Draught Delta (m)", "Score (%)", "Risk Level"]
    ws2.append(headers2)
    for col_num in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill

    for v in vessels:
        ws2.append([
            v.get("mmsi"),
            v.get("vessel_name"),
            v.get("flag"),
            v.get("imo"),
            v.get("vessel_type"),
            v.get("speed_knots"),
            v.get("course_deg"),
            v.get("distance_to_origin_km"),
            v.get("time_offset_hours"),
            v.get("draught_change_m"),
            v.get("compatibility_score"),
            v.get("risk_level")
        ])

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Sheet 3: Evidence Manifest
    ws3 = wb.create_sheet(title="Evidence Manifest")
    headers3 = ["Artifact ID", "Type", "Title", "SHA-256 Hash", "Source System", "Ingestion Timestamp", "Verified"]
    ws3.append(headers3)
    for col_num in range(1, len(headers3) + 1):
        cell = ws3.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = PatternFill(start_color="FF7A2F", end_color="FF7A2F", fill_type="solid")

    for ev in evidence:
        ws3.append([
            ev.get("id"),
            ev.get("artifact_type"),
            ev.get("title"),
            ev.get("sha256_hash"),
            ev.get("source_system"),
            ev.get("ingested_at"),
            "VERIFIED" if ev.get("verified") else "PENDING"
        ])

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(file_path)
    return file_path

def generate_geojson_report(incident: dict, detection: dict, drift: dict, vessels: list) -> str:
    ensure_reports_dir()
    inc_id = incident["id"]
    filename = f"SeaTrace_Geospatial_{inc_id.replace('-', '_')}.geojson"
    file_path = os.path.join(REPORTS_DIR, filename)

    features = []

    # Incident Point
    features.append({
        "type": "Feature",
        "geometry": {
            "type": "Point",
            "coordinates": [incident["lng"], incident["lat"]]
        },
        "properties": {
            "feature_type": "INCIDENT_CENTROID",
            "incident_id": inc_id,
            "title": incident.get("title"),
            "severity": incident.get("severity")
        }
    })

    # Detection polygon if available
    if detection and detection.get("mask_geojson"):
        try:
            poly_data = json.loads(detection["mask_geojson"])
            features.append({
                "type": "Feature",
                "geometry": poly_data.get("geometry", poly_data),
                "properties": {
                    "feature_type": "OBSERVED_SLICK_POLYGON",
                    "area_km2": detection.get("area_km2"),
                    "sensor": detection.get("sensor")
                }
            })
        except Exception:
            pass

    # Drift origin
    if drift and drift.get("origin_lat") and drift.get("origin_lng"):
        features.append({
            "type": "Feature",
            "geometry": {
                "type": "Point",
                "coordinates": [drift["origin_lng"], drift["origin_lat"]]
            },
            "properties": {
                "feature_type": "RECONSTRUCTED_RELEASE_ORIGIN",
                "uncertainty_km": drift.get("origin_uncertainty_km"),
                "time_window_hours": drift.get("time_window_hours")
            }
        })

    # Vessel tracks
    for v in vessels:
        if v.get("track_geojson"):
            try:
                track_data = json.loads(v["track_geojson"])
                features.append({
                    "type": "Feature",
                    "geometry": track_data,
                    "properties": {
                        "feature_type": "AIS_VESSEL_TRACK",
                        "mmsi": v.get("mmsi"),
                        "vessel_name": v.get("vessel_name"),
                        "score": v.get("compatibility_score"),
                        "risk_level": v.get("risk_level")
                    }
                })
            except Exception:
                pass

    fc = {
        "type": "FeatureCollection",
        "features": features,
        "metadata": {
            "generated_by": "SeaTrace AI Geospatial Engine",
            "incident_id": inc_id,
            "timestamp": datetime.utcnow().isoformat()
        }
    }

    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(fc, f, indent=2)

    return file_path
